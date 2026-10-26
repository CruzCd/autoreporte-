import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
import io
from datetime import datetime
import requests
import gc

def limpiar_datos(df):
    df = df.drop_duplicates()
    df = df.dropna(how='all')
    for col in df.select_dtypes(include='object').columns:
        df[col] = df[col].astype(str).str.strip()
    return df

def generar_resumen(df):
    resumen = {'filas': int(len(df)), 'columnas': int(len(df.columns)), 'numericas': []}
    for col in df.select_dtypes(include='number').columns:
        resumen['numericas'].append({
            'columna': str(col), 'suma': float(df[col].sum()),
            'promedio': float(df[col].mean()), 'maximo': float(df[col].max()),
            'minimo': float(df[col].min()),
        })
    return resumen

def generar_analisis_llm(resumen):
    datos_texto = f"Filas: {resumen['filas']}, Columnas: {resumen['columnas']}. "
    for item in resumen.get('numericas', []):
        datos_texto += f"Columna '{item['columna']}': Suma={item['suma']:.2f}, Promedio={item['promedio']:.2f}, Máximo={item['maximo']:.2f}, Mínimo={item['minimo']:.2f}. "

    prompt = f"""Eres un analista de datos experto. Basándote en este resumen de Excel, escribe un análisis conciso (máx. 3 párrafos) en español para un dueño de negocio. Destaca tendencias, valores atípicos y ofrece una recomendación práctica.

    Datos:
    {datos_texto}

    Análisis:"""

    try:
        # Endpoint público gratuito de Pollinations (No requiere API Key)
        response = requests.post(
            "https://text.pollinations.ai/openai",
            json={
                "model": "openai",
                "messages": [{"role": "user", "content": prompt}],
                "temperature": 0.5
            },
            timeout=30
        )
        response.raise_for_status()
        data = response.json()
        return data["choices"][0]["message"]["content"]
    except Exception as e:
        print(f"Error en Pollinations AI: {e}")
        return "No se pudo generar el análisis automático en este momento."

def generar_excel_reporte(df_limpio, resumen):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws['A1'] = "REPORTE AUTOMATIZADO"
    ws['A1'].font = Font(size=16, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells('A1:E1')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws['A2'] = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws['A2'].font = Font(italic=True, size=10)
    ws['A4'] = "Total de filas:"
    ws['B4'] = resumen['filas']
    ws['A5'] = "Total de columnas:"
    ws['B5'] = resumen['columnas']
    row = 7
    if resumen['numericas']:
        encabezados = ["Columna", "Suma", "Promedio", "Máximo", "Mínimo"]
        for i, texto in enumerate(encabezados, start=1):
            cell = ws.cell(row=row, column=i, value=texto)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        for item in resumen['numericas']:
            row += 1
            ws.cell(row=row, column=1, value=item['columna'])
            ws.cell(row=row, column=2, value=round(item['suma'], 2))
            ws.cell(row=row, column=3, value=round(item['promedio'], 2))
            ws.cell(row=row, column=4, value=round(item['maximo'], 2))
            ws.cell(row=row, column=5, value=round(item['minimo'], 2))
    for c in range(1, 6):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws2 = wb.create_sheet("Datos limpios")
    for c_idx, col_name in enumerate(df_limpio.columns, start=1):
        cell = ws2.cell(row=1, column=c_idx, value=col_name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for r_idx, row_data in enumerate(df_limpio.itertuples(index=False), start=2):
        for c_idx, value in enumerate(row_data, start=1):
            ws2.cell(row=r_idx, column=c_idx, value=value)
    for c in range(1, len(df_limpio.columns) + 1):
        ws2.column_dimensions[get_column_letter(c)].width = 18
    if resumen['numericas']:
        col_texto = df_limpio.select_dtypes(include='object').columns
        if len(col_texto) > 0:
            ws3 = wb.create_sheet("Gráfico")
            col_x = col_texto[0]
            col_y = resumen['numericas'][0]['columna']
            agrupado = df_limpio.groupby(col_x)[col_y].sum().reset_index().head(10)
            ws3.cell(row=1, column=1, value=col_x).font = Font(bold=True)
            ws3.cell(row=1, column=2, value=col_y).font = Font(bold=True)
            for i, fila in enumerate(agrupado.itertuples(index=False), start=2):
                ws3.cell(row=i, column=1, value=fila[0])
                ws3.cell(row=i, column=2, value=fila[1])
            chart = BarChart()
            chart.title = f"{col_y} por {col_x}"
            data = Reference(ws3, min_col=2, min_row=1, max_row=len(agrupado) + 1)
            cats = Reference(ws3, min_col=1, min_row=2, max_row=len(agrupado) + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws3.add_chart(chart, "D5")
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def procesar_archivo(file_stream, filename):
    if filename.lower().endswith('.csv'):
        df = pd.read_csv(file_stream)
    else:
        df = pd.read_excel(file_stream)
    df_limpio = limpiar_datos(df)
    resumen = generar_resumen(df_limpio)
    reporte = generar_excel_reporte(df_limpio, resumen)
    del df
    del df_limpio
    gc.collect()
    return reporte, resumen
